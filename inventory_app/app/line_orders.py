"""LINE 接單：匯入聊天匯出檔、只保存新留言、用 OpenAI 轉成訂單並對應店內商品。

OpenAI 金鑰：環境變數 OPENAI_API_KEY；本機版也可放在資料夾 data/openai-key.txt。
模型：OPENAI_MODEL（預設 gpt-6-luna），思考程度：OPENAI_REASONING_EFFORT（選填），
API 位址：OPENAI_BASE_URL（選填，預設 https://api.openai.com/v1）。
"""
import io
import json
import math
import os
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path

from flask import Blueprint, current_app, jsonify, request, session

from . import line_chat as lc
from .catalog_order import product_key
from .common import Problem, decimal_text, now, number, paginate, required_text, today, transaction
from .dashboard import products_with_stock
from .db import get_db

line_orders = Blueprint('line_orders', __name__)
KEY_FILE = 'openai-key.txt'
DEFAULT_MODEL = 'gpt-6-luna'
MAX_EXPORT_BYTES = 30 * 1024 * 1024
STALE_CONVERSION = timedelta(minutes=10)
ACTIONS = ('新訂單', '追加', '修改', '取消')
STATUSES = ('open', 'done', 'dismissed')
SORTS = {'new': 'DESC', 'old': 'ASC'}  # 依留言時間：新到舊、舊到新


def ai_settings():
    model = os.environ.get('OPENAI_MODEL') or current_app.config.get('OPENAI_MODEL') or DEFAULT_MODEL
    effort = os.environ.get('OPENAI_REASONING_EFFORT') or current_app.config.get('OPENAI_REASONING_EFFORT')
    return model, effort or None


def key_file():
    return Path(current_app.config['DATA_DIR']) / KEY_FILE


def api_key():
    key = os.environ.get('OPENAI_API_KEY', '').strip() or current_app.config.get('OPENAI_API_KEY', '')
    if key or current_app.config.get('CLOUD_MODE'):
        return key
    path = key_file()
    if path.is_file():
        for line in path.read_text('utf-8-sig').splitlines():
            line = line.strip()
            if line and not line.startswith('#'):
                return line.removeprefix('OPENAI_API_KEY=').strip().strip('"\'')
    return ''


def base_url():
    return (os.environ.get('OPENAI_BASE_URL') or 'https://api.openai.com/v1').rstrip('/')


def transport():
    # 測試用假的 OpenAI；正式環境需要金鑰。
    fake = current_app.config.get('OPENAI_TRANSPORT')
    if fake:
        return fake
    key = api_key()
    return lc.http_send(key, base_url() + '/responses') if key else None


def text(value, limit):
    return str(value).strip()[:limit] if value is not None else ''


def amount(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
        return None
    return decimal_text(round(Decimal(str(value)), 6))


def parse_start(value):
    value = (value or '').strip()
    if not value:
        return None
    if value == 'baseline':
        return value
    try:
        return lc.parse_since(value)
    except ValueError as e:
        raise Problem(str(e)) from None


def import_export(conn, filename, data, actor, start=None):
    """把整份匯出檔和上次位置比對，只存新留言。start：None、'baseline' 或 (日期, 時間)。"""
    try:
        messages = lc.parse_chat(lc.decode_export(data))
    except ValueError as e:
        raise Problem(str(e)) from None
    if not messages:
        raise Problem('這個檔案找不到 LINE 留言，請確認是 LINE「儲存聊天」匯出的 .txt')
    named = lc.chat_name(filename)
    last = messages[-1]
    with transaction(conn):
        # 依內容找出是哪個群組，檔名可能被改過，例如「…(更新1).txt」；同名的先試。
        states = sorted((dict(r) for r in conn.execute('SELECT * FROM line_chats')), key=lambda s: s['chat'] != named)
        state = found = None
        for candidate in states:
            found = lc.find_new_messages(json.loads(candidate['tail']), messages)
            if found is not None:
                state = candidate
                break
        if state is None:
            state = next((s for s in states if s['chat'] == named), None)
        chat = state['chat'] if state else named
        info = {'chat': chat, 'last_on': last.date, 'last_at': last.time, 'message_count': len(messages)}
        missing = 0
        if start is None:
            if not state:
                raise Problem(f'第一次匯入「{chat}」，請選擇從哪個時間開始接單',
                              status=409, code='line_first_import', fields=info)
            if last.key < (state['last_on'], state['last_at']):
                raise Problem(f"這份匯出檔最後一則是 {last.date} {last.time}，比上次匯入的 "
                              f"{state['last_on']} {state['last_at']} 還舊，請重新從 LINE 儲存聊天",
                              status=409, code='line_older_export')
            if found is None:
                raise Problem('在這份匯出檔找不到上次匯入的位置（可能是不同群組，或很多留言被收回）。'
                              '請確認檔案；若沒問題，可重新選擇開始時間。',
                              status=409, code='line_position_lost',
                              fields={**info, 'previous_on': state['last_on'], 'previous_at': state['last_at']})
            new, missing = found
        else:
            new = [m for m in messages if start != 'baseline' and m.key >= start]
            if new:
                # 重新設定起點時，略過已匯入過的留言，避免重複訂單。
                seen = {(r['sent_on'], r['sent_at'], r['raw']) for r in conn.execute(
                    'SELECT sent_on,sent_at,raw FROM line_messages WHERE sent_on>=?', (new[0].date,))}
                new = [m for m in new if (m.date, m.time, m.raw) not in seen]
        menu = lc.latest_menu(messages)
        if menu is None and state and state['menu']:
            menu = json.loads(state['menu'])
        values = (last.date, last.time, len(messages), json.dumps([m.fp for m in messages[-lc.TAIL_SIZE:]]),
                  json.dumps(menu, ensure_ascii=False) if menu else None, now())
        if state:
            conn.execute('UPDATE line_chats SET last_on=?,last_at=?,message_count=?,tail=?,menu=?,updated_at=? '
                         'WHERE chat=?', values + (chat,))
        else:
            conn.execute('INSERT INTO line_chats(chat,last_on,last_at,message_count,tail,menu,updated_at) '
                         'VALUES(?,?,?,?,?,?,?)', (chat,) + values)
        if not new and start is None:
            return {**info, 'import_id': None, 'new_messages': 0, 'missing_messages': missing}
        kinds = [lc.classify(m) for m in new]
        customers = kinds.count('customer')
        status = 'baseline' if not new else ('pending' if customers else 'empty')
        iid = conn.execute('''INSERT INTO line_imports(chat,filename,created_at,actor,total_messages,new_messages,
            customer_messages,missing_messages,status) VALUES(?,?,?,?,?,?,?,?,?)''',
            (chat, filename[:200], now(), str(actor), len(messages), len(new), customers, missing, status)).lastrowid
        for m, kind in zip(new, kinds):
            conn.execute('INSERT INTO line_messages(import_id,sent_on,sent_at,kind,raw) VALUES(?,?,?,?,?)',
                         (iid, m.date, m.time, kind, m.raw))
    return {**info, 'import_id': iid, 'new_messages': len(new), 'customer_messages': customers,
            'missing_messages': missing, 'status': status}


def normalize_orders(data, numbers, product_ids):
    """AI 回覆 → 要存的訂單；來源留言與時間由程式填入，不採用 AI 抄寫的內容。"""
    result = []
    for o in data.get('orders') or []:
        if not isinstance(o, dict):
            continue
        ids = sorted({i for i in o.get('source_ids') or [] if type(i) is int and i in numbers})
        rows = [numbers[i] for i in ids]
        first = rows[0] if rows else numbers[min(numbers)]
        reason = text(o.get('review_reason'), 200)
        items = []
        for item in (o.get('items') or [])[:100]:
            if not isinstance(item, dict):
                continue
            pid = item.get('product_id')
            items.append({'name': text(item.get('name'), 100) or '（未寫品名）',
                          'product_id': pid if type(pid) is int and pid in product_ids else None,
                          'quantity': amount(item.get('quantity')), 'unit': text(item.get('unit'), 20),
                          'unit_price': amount(item.get('unit_price')),
                          'processing': text(item.get('processing'), 100), 'note': text(item.get('note'), 200)})
        result.append({'sent_on': first['sent_on'], 'sent_at': first['sent_at'],
                       'customer': text(o.get('customer'), 100), 'location': text(o.get('location'), 200),
                       'action': o.get('action') if o.get('action') in ACTIONS else '新訂單',
                       'carrier': text(o.get('carrier'), 60), 'payment': text(o.get('payment'), 200),
                       'note': text(o.get('note'), 500),
                       'needs_review': bool(o.get('needs_review')) or not rows,
                       'review_reason': reason or ('' if rows else 'AI 沒有標明來源留言，請核對原文'),
                       'source_ids': [r['id'] for r in rows], 'items': items})
    return result


def stale(started_at):
    try:
        return datetime.fromisoformat(started_at) < datetime.fromisoformat(now()) - STALE_CONVERSION
    except (TypeError, ValueError):
        return True


def convert_import(conn, iid):
    send = transport()
    if send is None:
        raise Problem('尚未設定 OpenAI 金鑰，新留言已保存；設定金鑰後按「重新轉換」',
                      status=409, code='line_no_key')
    model, effort = ai_settings()
    started = now()
    with transaction(conn):
        imp = conn.execute('SELECT * FROM line_imports WHERE id=?', (iid,)).fetchone()
        if not imp:
            raise Problem('找不到這次匯入', status=404)
        if imp['status'] == 'converting' and not stale(imp['started_at']):
            raise Problem('這批正在轉換中，請稍候再重新整理', status=409)
        if imp['status'] not in ('pending', 'failed', 'converting'):
            raise Problem('這批不需要轉換', status=409)
        conn.execute("UPDATE line_imports SET status='converting',started_at=?,error=NULL WHERE id=?", (started, iid))
        rows = [dict(r) for r in conn.execute(
            "SELECT * FROM line_messages WHERE import_id=? AND kind='customer' ORDER BY id", (iid,))]
        state = conn.execute('SELECT menu FROM line_chats WHERE chat=?', (imp['chat'],)).fetchone()
        products = [dict(r) for r in conn.execute('SELECT id,name,unit,price FROM products ORDER BY category,code,id')]
    menu = json.loads(state['menu'])['text'] if state and state['menu'] else ''
    candidates = [lc.Message(r['sent_on'], r['sent_at'], r['raw']) for r in rows]
    catalog, product_ids = lc.products_text(products), {p['id'] for p in products}
    usage, orders = {}, []
    try:
        for chunk in lc.chunk_blocks(lc.numbered_blocks(candidates)):
            payload = lc.build_payload(model, catalog, menu, chunk, effort)
            data, used = lc.read_response(lc.send_with_retry(send, payload))
            lc.add_usage(usage, used)
            orders += normalize_orders(data, {i: rows[i - 1] for i, _ in chunk}, product_ids)
    except Exception as e:
        message = str(e) if isinstance(e, lc.AIError) else 'AI 轉換時發生錯誤'
        if not isinstance(e, lc.AIError):
            current_app.logger.exception('LINE order conversion failed')
        with transaction(conn):
            conn.execute("UPDATE line_imports SET status='failed',error=? WHERE id=? AND status='converting' "
                         'AND started_at=?', (message, iid, started))
        raise Problem(message + '。新留言已保存，可稍後按「重新轉換」', status=502, code='line_ai_failed') from None
    cost = lc.estimate_cost(model, usage)
    with transaction(conn):
        claimed = conn.execute("""UPDATE line_imports SET status='done',error=NULL,model=?,usage=?,cost=?,converted_at=?
            WHERE id=? AND status='converting' AND started_at=?""",
            (model, json.dumps(usage), None if cost is None else f'{cost:.6f}', now(), iid, started)).rowcount
        if claimed != 1:
            raise Problem('這批已在其他畫面轉換，請重新整理', status=409)
        at = now()
        for o in orders:
            oid = conn.execute('''INSERT INTO line_orders(import_id,sent_on,sent_at,customer,location,action,carrier,
                payment,note,needs_review,review_reason,status,source_ids,created_at,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,'open',?,?,?)''',
                (iid, o['sent_on'], o['sent_at'], o['customer'], o['location'], o['action'], o['carrier'],
                 o['payment'], o['note'], int(o['needs_review']), o['review_reason'],
                 json.dumps(o['source_ids']), at, at)).lastrowid
            for item in o['items']:
                conn.execute('''INSERT INTO line_order_items(order_id,name,product_id,quantity,unit,unit_price,
                    processing,note) VALUES(?,?,?,?,?,?,?,?)''',
                    (oid, item['name'], item['product_id'], item['quantity'], item['unit'], item['unit_price'],
                     item['processing'], item['note']))
    return len(orders)


def import_row(row):
    result = dict(row)
    result['usage'] = json.loads(row['usage']) if row['usage'] else None
    return result


def import_detail(conn, iid):
    row = conn.execute('SELECT * FROM line_imports WHERE id=?', (iid,)).fetchone()
    if not row:
        raise Problem('找不到這次匯入', status=404)
    result = import_row(row)
    messages = [dict(r) for r in conn.execute('SELECT * FROM line_messages WHERE import_id=? ORDER BY id', (iid,))]
    used = set()
    for r in conn.execute('SELECT source_ids FROM line_orders WHERE import_id=?', (iid,)):
        used.update(json.loads(r['source_ids']))
    pick = lambda kind: [{k: m[k] for k in ('sent_on', 'sent_at', 'raw')} for m in messages if m['kind'] == kind
                         and (kind != 'customer' or m['id'] not in used)]
    result['orders'] = conn.execute('SELECT COUNT(*) FROM line_orders WHERE import_id=?', (iid,)).fetchone()[0]
    result['not_orders'] = pick('customer') if row['status'] == 'done' else []
    result['images'], result['recalls'] = pick('image'), pick('recall')
    result['staff_messages'] = sum(m['kind'] == 'staff' for m in messages)
    return result


def ai_status():
    """畫面只顯示有沒有接上；模型與費用留在匯入紀錄供查核。"""
    cloud = bool(current_app.config.get('CLOUD_MODE'))
    return {'configured': bool(current_app.config.get('OPENAI_TRANSPORT') or api_key()), 'cloud': cloud,
            'key_location': 'Google Cloud Secret Manager 的 OPENAI_API_KEY' if cloud else str(key_file())}


def time_range(args):
    """from／to 可填「2026-10-03T09:00」或日期；只填日期的結束時間算到 23:59。都沒填代表全部時間。"""
    raw_from, raw_to = (args.get('from') or '').strip(), (args.get('to') or '').strip()
    if not raw_from and not raw_to:
        return None
    try:
        start = lc.parse_since(raw_from) if raw_from else ('0000-01-01', '00:00')
        end = lc.parse_since(raw_to) if raw_to else ('9999-12-31', '23:59')
    except ValueError as e:
        raise Problem(str(e)) from None
    if len(raw_to) == 10:
        end = (end[0], '23:59')
    if start > end:
        raise Problem('結束時間要晚於開始時間')
    return f'{start[0]} {start[1]}', f'{end[0]} {end[1]}'


PERIOD_SQL = "sent_on || ' ' || sent_at BETWEEN ? AND ?"


def rows_by_ids(conn, query, ids):
    """分批查詢，避免 IN 參數超過資料庫上限。"""
    for start in range(0, len(ids), 500):
        part = ids[start:start + 500]
        yield from conn.execute(query.format(','.join('?' for _ in part)), part)


def sale_links(conn, ids):
    """訂單編號 → 已扣庫存的每日銷售單編號；沖銷掉的銷售單不算。"""
    return {r['order_id']: r['sheet_id'] for r in rows_by_ids(conn, """SELECT s.order_id,s.sheet_id
        FROM line_order_sales s JOIN daily_sales d ON d.id=s.sheet_id
        WHERE d.status='posted' AND s.order_id IN ({})""", list(ids))}


def attach_details(conn, orders):
    if not orders:
        return orders
    items = {}
    for r in rows_by_ids(conn, 'SELECT * FROM line_order_items WHERE order_id IN ({}) ORDER BY id',
                         [o['id'] for o in orders]):
        items.setdefault(r['order_id'], []).append(dict(r))
    stock = {p['id']: p for p in products_with_stock(conn, today())}
    sales = sale_links(conn, [o['id'] for o in orders])
    source_ids = sorted({i for o in orders for i in json.loads(o['source_ids'])})
    sources = {r['id']: dict(r) for r in rows_by_ids(
        conn, 'SELECT id,sent_on,sent_at,raw FROM line_messages WHERE id IN ({})', source_ids)}
    for o in orders:
        o['needs_review'] = bool(o['needs_review'])
        o['sale_id'] = sales.get(o['id'])
        o['items'] = items.get(o['id'], [])
        for item in o['items']:
            p = stock.get(item['product_id'])
            item['product'] = {k: p[k] for k in ('id', 'code', 'name', 'unit', 'price', 'quantity', 'image')} if p else None
        o['source'] = [sources[i] for i in json.loads(o.pop('source_ids')) if i in sources]
    return orders


def prep_summary(orders):
    """商品數量合計：同商品同單位加總；沒寫數量、需確認的另外計次。不是訂單與取消不計。"""
    groups = {}
    for o in orders:
        if o['status'] == 'dismissed' or o['action'] == '取消':
            continue
        for item in o['items']:
            p = item['product']
            key = (('p', p['id']) if p else ('n', item['name'])) + (item['unit'],)
            g = groups.setdefault(key, {'name': p['name'] if p else item['name'], 'unit': item['unit'],
                                        'product_id': p['id'] if p else None, 'code': p['code'] if p else '',
                                        'total': Decimal(0), 'orders': 0, 'unknown': 0, 'review': 0})
            g['orders'] += 1
            g['review'] += int(o['needs_review'])
            if item['quantity'] is None:
                g['unknown'] += 1
            else:
                g['total'] += Decimal(item['quantity'])
    rows = [{**g, 'total': decimal_text(g['total'])} for g in groups.values()]
    return sorted(rows, key=lambda g: (g['product_id'] is None, -g['orders'], g['name']))


@line_orders.get('/api/line-orders')
def list_orders():
    status, sort = request.args.get('status', 'open'), request.args.get('sort', 'new')
    if status not in STATUSES:
        raise Problem('篩選狀態不正確')
    if sort not in SORTS:
        raise Problem('排序方式不正確')
    period = time_range(request.args)
    conn = get_db()
    latest = conn.execute('SELECT MAX(id) AS id FROM line_imports').fetchone()['id']
    where, params = ['status=?'], [status]
    if period:
        where.append(PERIOD_SQL)
        params += period
    # 全部依留言時間排；最新一次匯入的訂單另外標「新」，不再排到最前面。
    order = ' ORDER BY sent_on {0},sent_at {0},id {0}'.format(SORTS[sort])
    rows = [dict(r) for r in conn.execute('SELECT * FROM line_orders WHERE ' + ' AND '.join(where) + order, params)]
    for row in rows:
        row['is_new'] = row['import_id'] == latest
    if status == 'open' or period:  # 右側合計＝清單上全部訂單的加總
        page = paginate(attach_details(conn, rows), request.args.get('page', 1), 30, maximum=30)
        summary = prep_summary(rows)
    else:
        page = paginate(rows, request.args.get('page', 1), 30, maximum=30)
        attach_details(conn, page['items'])
        summary = None
    counts = {s: 0 for s in STATUSES}
    for r in conn.execute('SELECT status,COUNT(*) AS n FROM line_orders' + (' WHERE ' + PERIOD_SQL if period else '')
                          + ' GROUP BY status', period or ()):
        counts[r['status']] = r['n']
    imports = [import_row(r) for r in conn.execute(
        'SELECT i.*,(SELECT COUNT(*) FROM line_orders o WHERE o.import_id=i.id) AS order_count '
        'FROM line_imports i ORDER BY i.id DESC LIMIT 8')]
    return jsonify({**page, 'counts': counts, 'summary': summary, 'new_count': sum(r['is_new'] for r in rows),
                    'imports': imports, 'latest': import_detail(conn, latest) if latest else None, 'ai': ai_status()})


def orders_workbook(title, info, summary, orders, others):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows, widths):
        ws = wb.create_sheet(name)
        ws.append([title])
        ws.append([info])
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # 客人打的字一律當文字，「=」開頭也不會變成公式
        ws['A1'].font = Font(size=16, bold=True, color='435B32')
        ws['A2'].font = Font(color='6A705F')
        for cell in ws[3]:
            cell.fill = PatternFill('solid', fgColor='435B32')
            cell.font = Font(color='FFFFFF', bold=True)
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = 'A4'
        ws.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{max(3, ws.max_row)}'

    number = lambda v: float(v) if v is not None else None
    sheet('商品合計', ['商品', '數量合計', '單位', '訂單數', '沒寫數量', '需確認', '店內品號'],
          [[g['name'], number(g['total']) if g['total'] != '0' or not g['unknown'] else None, g['unit'],
            g['orders'], g['unknown'] or None, g['review'] or None, g['code']] for g in summary],
          [32, 12, 8, 10, 10, 10, 12])
    names = {'open': '待處理', 'done': '已完成'}
    rows = []
    for o in orders:
        raw = '\n\n'.join(m['raw'] for m in o['source'])
        review = (o['review_reason'] or '需確認') if o['needs_review'] else ''
        # 店內品號、數量可以在 Excel 裡改，再到「每日銷售單」匯入；訂單編號用來避免重複扣庫存。
        for n, item in enumerate(o['items'] or [None]):
            body = ([item['name'], number(item['quantity']), item['unit'],
                     item['product']['code'] if item['product'] else '', item['processing'], item['note']]
                    if item else ['（沒有辨識出品項）', None, '', '', '', ''])
            rows.append([o['sent_on'], o['sent_at'], o['location'], o['customer'], o['action']] + body
                        + [o['carrier'], o['payment'], o['note'], names.get(o['status'], o['status']), review,
                           raw if n == 0 else '', o['id'], f"#{o['sale_id']}" if o.get('sale_id') else ''])
    sheet('訂單明細', ['日期', '時間', '地點', '客人', '類型', '商品', '數量', '單位', '店內品號', '處理方式', '品項備註',
                       '載具', '付款', '訂單備註', '狀態', '需確認', 'LINE 原文', '訂單編號', '已扣庫存銷售單'],
          rows, [11, 7, 24, 16, 8, 26, 7, 6, 10, 14, 18, 12, 16, 22, 8, 20, 50, 9, 12])
    sheet('其他留言', ['日期', '時間', '類型', 'LINE 原文'],
          [[m['sent_on'], m['sent_at'], m['label'], m['raw']] for m in others], [11, 7, 26, 80])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def period_label(period):
    if not period:
        return '全部時間'
    if period[0][:10] == period[1][:10] and period[0][11:] == '00:00' and period[1][11:] == '23:59':
        return period[0][:10]
    return f'{period[0]} ～ {period[1]}'


@line_orders.get('/api/line-orders/export')
def export_orders():
    """選定時間內「待處理＋已完成」的訂單：商品合計、訂單明細（含原文）、沒轉成訂單的留言。"""
    from flask import send_file
    period = time_range(request.args)
    conn = get_db()
    clause, params = (' AND ' + PERIOD_SQL, period) if period else ('', ())
    orders = [dict(r) for r in conn.execute("SELECT * FROM line_orders WHERE status IN ('open','done')" + clause
                                            + ' ORDER BY sent_on,sent_at,id', params)]
    if not orders:
        raise Problem('這段時間沒有訂單', status=404)
    attach_details(conn, orders)
    linked, dismissed = set(), set()
    for r in conn.execute('SELECT status,source_ids FROM line_orders' + (' WHERE ' + PERIOD_SQL if period else ''),
                          params):
        (dismissed if r['status'] == 'dismissed' else linked).update(json.loads(r['source_ids']))
    kinds = {'image': '客人傳圖片（請到 LINE 查看）', 'recall': '收回的訊息'}
    others = []
    for m in conn.execute("SELECT * FROM line_messages WHERE kind IN ('customer','image','recall')" + clause
                          + ' ORDER BY sent_on,sent_at,id', params):
        if m['kind'] == 'customer' and m['id'] in linked:
            continue
        label = kinds.get(m['kind']) or ('已標記不是訂單' if m['id'] in dismissed else 'AI 沒當成訂單，請看一下')
        others.append({'sent_on': m['sent_on'], 'sent_at': m['sent_at'], 'label': label, 'raw': m['raw']})
    label = period_label(period)
    done = sum(o['status'] == 'done' for o in orders)
    info = f"匯出時間 {now()[:16].replace('T', ' ')}｜共 {len(orders)} 筆訂單（待處理 {len(orders) - done}、已完成 {done}）"
    book = orders_workbook(f'LINE 訂單｜{label}', info, prep_summary(orders), orders, others)
    name = 'LINE訂單_' + label.replace(' ～ ', '_至_').replace(':', '').replace(' ', '_') + '.xlsx'
    return send_file(book, as_attachment=True, download_name=name, max_age=0,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@line_orders.post('/api/line-orders/import')
def upload_export():
    upload = request.files.get('file')
    if not upload or not upload.filename:
        raise Problem('請選擇 LINE 匯出的 .txt 檔')
    filename = Path(upload.filename.replace('\\', '/')).name
    if not filename.lower().endswith('.txt'):
        raise Problem('請選擇 .txt 檔（LINE「儲存聊天」匯出的文字檔）')
    data = upload.read(MAX_EXPORT_BYTES + 1)
    if len(data) > MAX_EXPORT_BYTES:
        raise Problem('檔案超過 30 MB，不像 LINE 聊天匯出檔')
    conn = get_db()
    result = import_export(conn, filename, data, session['user'], parse_start(request.form.get('start')))
    result['orders_created'] = 0
    if result['import_id'] and result['status'] == 'pending':
        try:
            result['orders_created'] = convert_import(conn, result['import_id'])
        except Problem as e:
            result['convert_error'] = e.message
    if result['import_id']:
        result['import'] = import_detail(conn, result['import_id'])
        result['status'] = result['import']['status']
    return jsonify(result)


@line_orders.get('/api/line-orders/imports/<int:iid>')
def get_import(iid):
    return jsonify(import_detail(get_db(), iid))


@line_orders.post('/api/line-orders/imports/<int:iid>/convert')
def retry_import(iid):
    conn = get_db()
    created = convert_import(conn, iid)
    return jsonify({'orders_created': created, 'import': import_detail(conn, iid)})


def optional_text(value, label, limit):
    if value is None:
        return ''
    if not isinstance(value, str) or len(value.strip()) > limit:
        raise Problem(f'{label}最多 {limit} 字')
    return value.strip()


def edited_items(raw_items, product_ids):
    if not isinstance(raw_items, list) or len(raw_items) > 100:
        raise Problem('品項最多 100 項')
    items = []
    for n, raw in enumerate(raw_items, 1):
        try:
            if not isinstance(raw, dict):
                raise Problem('品項格式不正確')
            pid = raw.get('product_id')
            if pid is not None and (type(pid) is not int or pid not in product_ids):
                raise Problem('店內商品不存在，請重新選擇')
            items.append((required_text(raw.get('name'), '品名', 100), pid,
                          decimal_text(number(raw.get('quantity'), '數量', optional=True, positive=True)),
                          optional_text(raw.get('unit'), '單位', 20),
                          decimal_text(number(raw.get('unit_price'), '單價', optional=True)),
                          optional_text(raw.get('processing'), '處理方式', 100),
                          optional_text(raw.get('note'), '品項備註', 200)))
        except Problem as error:
            raise Problem(f'第 {n} 項：{error.message}') from None
    return items


@line_orders.get('/api/line-orders/products')
def order_products():
    """修改訂單時選「店內商品」用。"""
    rows = [dict(r) for r in get_db().execute('SELECT id,code,name,unit,category FROM products')]
    return jsonify(items=sorted(rows, key=product_key))


@line_orders.put('/api/line-orders/<int:oid>')
def edit_order(oid):
    """老闆修正 AI 整理的訂單：整筆換成畫面上的內容。已扣庫存的訂單要先沖銷銷售單才能改。"""
    body = request.get_json(silent=True)
    if not isinstance(body, dict) or type(body.get('version')) is not int:
        raise Problem('訂單內容或版本不正確')
    if body.get('action') not in ACTIONS:
        raise Problem('訂單類型不正確')
    if type(body.get('needs_review')) is not bool:
        raise Problem('「需確認」格式不正確')
    fields = {key: optional_text(body.get(key), label, limit) for key, label, limit in (
        ('location', '地點', 100), ('customer', '客人', 100), ('carrier', '載具', 100),
        ('payment', '付款', 100), ('note', '備註', 500))}
    reason = optional_text(body.get('review_reason'), '需確認原因', 200) if body['needs_review'] else ''
    conn = get_db()
    items = edited_items(body.get('items'), {r['id'] for r in conn.execute('SELECT id FROM products')})
    with transaction(conn):
        sheet = sale_links(conn, [oid]).get(oid)
        if sheet:
            raise Problem(f'這筆已帶入每日銷售單 #{sheet} 扣庫存；要修改請先到每日銷售單沖銷那張單', status=409)
        stamp = now()
        changed = conn.execute('''UPDATE line_orders SET location=?,customer=?,action=?,carrier=?,payment=?,note=?,
            needs_review=?,review_reason=?,edited_at=?,updated_at=?,version=version+1 WHERE id=? AND version=?''',
            (fields['location'], fields['customer'], body['action'], fields['carrier'], fields['payment'],
             fields['note'], int(body['needs_review']), reason, stamp, stamp, oid, body['version'])).rowcount
        if changed != 1:
            if conn.execute('SELECT 1 FROM line_orders WHERE id=?', (oid,)).fetchone():
                raise Problem('這筆訂單剛被更新，請重新整理', status=409)
            raise Problem('找不到訂單', status=404)
        conn.execute('DELETE FROM line_order_items WHERE order_id=?', (oid,))
        for item in items:
            conn.execute('''INSERT INTO line_order_items(order_id,name,product_id,quantity,unit,unit_price,processing,note)
                VALUES(?,?,?,?,?,?,?,?)''', (oid,) + item)
        row = dict(conn.execute('SELECT * FROM line_orders WHERE id=?', (oid,)).fetchone())
    return jsonify(attach_details(conn, [row])[0])


@line_orders.patch('/api/line-orders/<int:oid>')
def update_order(oid):
    body = request.get_json()
    status, version = body.get('status'), body.get('version')
    if status not in STATUSES or type(version) is not int:
        raise Problem('訂單狀態或版本不正確')
    conn = get_db()
    with transaction(conn):
        sheet = sale_links(conn, [oid]).get(oid)
        if sheet:
            raise Problem(f'這筆已帶入每日銷售單 #{sheet} 扣庫存；要修改請先到每日銷售單沖銷那張單', status=409)
        changed = conn.execute('UPDATE line_orders SET status=?,version=version+1,updated_at=? WHERE id=? AND version=?',
                               (status, now(), oid, version)).rowcount
        if changed != 1:
            if conn.execute('SELECT 1 FROM line_orders WHERE id=?', (oid,)).fetchone():
                raise Problem('這筆訂單剛被更新，請重新整理', status=409)
            raise Problem('找不到訂單', status=404)
        row = conn.execute('SELECT id,status,version FROM line_orders WHERE id=?', (oid,)).fetchone()
    return jsonify(dict(row))


@line_orders.post('/api/line-orders/check-key')
def check_key():
    """用「查詢模型」確認金鑰與模型可用；不產生費用。"""
    import requests
    key = api_key()
    if not key:
        raise Problem('尚未設定 OpenAI 金鑰', status=409, code='line_no_key')
    model, _ = ai_settings()
    try:
        response = requests.get(f'{base_url()}/models/{model}', timeout=20,
                                headers={'Authorization': f'Bearer {key}'})
    except requests.RequestException:
        raise Problem('連不上 OpenAI，請檢查網路', status=502) from None
    if response.status_code == 200:
        return jsonify(ok=True, model=model, message=f'連線成功，可以使用 {model}')
    try:
        error = response.json().get('error') or {}
    except ValueError:
        error = {}
    raise Problem(str(lc.ai_error(response.status_code, error, model)), status=502)
