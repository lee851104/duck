"""Read standalone sales workbooks for human review; never mutate inventory."""
import hashlib
import io
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .common import Problem, decimal_text, valid_date

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
NO_ITEMS = '（沒有辨識出品項）'
ALIASES = {
    'code': ('店內品號', '品號', '商品編號'),
    'name': ('商品', '商品名稱', '品名'),
    'quantity': ('數量', '銷售數量', '賣出數量'),
    'unit': ('單位',),
}
# 舊 LINE 接單匯出與本系統銷售明細的狀態欄：避開取消單、已扣過庫存的列，其餘標示給老闆核對。
OPTIONAL = {
    'kind': ('類型',),
    'status': ('狀態',),
    'sheet': ('已扣庫存銷售單', '銷售單號'),
    'sold_on': ('銷售日期',),
}


def text(value):
    return '' if value is None else str(value).strip()


def column(names, aliases):
    return next((names.index(alias) for alias in aliases if alias in names), None)


def day_text(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    try:
        return date.fromisoformat(text(value)[:10].replace('/', '-')).isoformat()
    except ValueError:
        return None


def preview_excel(conn, stream, filename, sold_on=None):
    day = valid_date(sold_on) if sold_on else None
    data = stream.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise Problem('Excel 不可超過 10 MB，請分批匯入')
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, ValueError, OSError):
        raise Problem('無法讀取 Excel，請確認是有效的 .xlsx 檔') from None
    products = [dict(p) for p in conn.execute('SELECT id,code,name,unit FROM products')]
    by_code = {p['code']: p for p in products}
    by_name = defaultdict(list)
    for p in products:
        by_name[p['name']].append(p)
    sheets = {}

    def sheet_info(sid):
        if sid not in sheets:
            row = conn.execute('SELECT id,status,sold_on FROM daily_sales WHERE id=?', (sid,)).fetchone()
            sheets[sid] = dict(row) if row else None
        return sheets[sid]

    items, skipped = [], []
    try:
        # Prefer detail sheets so an export's summary is never imported twice.
        sheet = next((book[n] for n in ('銷售明細', '訂單明細') if n in book.sheetnames), book.worksheets[0])
        columns = None
        for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
            if number > MAX_ROWS:
                raise Problem('Excel 超過 5000 列，請分批匯入')
            if columns is None:
                names = [text(v) for v in values]
                found = {key: column(names, aliases) for key, aliases in ALIASES.items()}
                if found['quantity'] is not None and (found['code'] is not None or found['name'] is not None):
                    columns = {**found, **{key: column(names, aliases) for key, aliases in OPTIONAL.items()}}
                elif number >= 10:
                    break
                continue
            if not any(v not in (None, '') for v in values):
                continue
            row = {key: values[index] if index is not None and index < len(values) else None
                   for key, index in columns.items()}
            code, name = text(row['code']), text(row['name'])
            if name == NO_ITEMS:
                continue
            label = f'Excel 第 {number} 列 · {name or code or "未填品名"}'
            kind, status, row_day = text(row['kind']), text(row['status']), day_text(row['sold_on'])
            if kind == '取消':
                skipped.append({'id': None, 'kind': 'cancel', 'label': label, 'reason': '取消單，不扣庫存'})
                continue
            warning = []
            reference = text(row['sheet'])
            if reference:
                matched = re.fullmatch(r'#?\s*(\d{1,12})(?:\.0+)?', reference)
                recorded = sheet_info(int(matched[1])) if matched else None
                if recorded and recorded['status'] == 'posted' and row_day in (None, recorded['sold_on']):
                    skipped.append({'id': None, 'kind': 'linked', 'label': label,
                                    'reason': f"已在每日銷售單 #{recorded['id']}（{recorded['sold_on']}）扣過庫存，"
                                              '不會重複扣；要重填請先沖銷原單'})
                    continue
                # 已沖銷的單可以重新帶入；找不到或日期不同的單無法確認，交給老闆核對。
                if not recorded or row_day not in (None, recorded['sold_on']):
                    warning.append(f'檔案標示已在銷售單 {reference} 扣過庫存，但這台電腦找不到相同的單，請確認沒有重複扣')
            elif status == '已扣庫存':
                warning.append('這列標示「已扣庫存」，請確認沒有重複扣')
            if kind == '修改':
                warning.append('修改單：請確認沒有和原本的訂單重複扣')
            if status == '待處理':
                warning.append('LINE 接單還是「待處理」，請確認已出貨')
            if day and row_day and row_day != day:
                warning.append(f'這列的銷售日期是 {row_day}，和這張單的 {day} 不同')
            product = by_code.get(code) if code else None
            if code and not product:
                warning.append(f'找不到品號「{code}」，請手動選商品')
            elif not code and len(by_name[name]) == 1:
                product = by_name[name][0]
            elif not code:
                warning.append('品名無法唯一對應，請手動選商品')
            if product and name and name != product['name']:
                warning.append('品號與品名不一致，請核對對應商品')
            quantity = None
            try:
                value = Decimal(text(row['quantity']))
                if value.is_finite() and 0 < value <= 999999999 and value == value.quantize(Decimal('0.000001')):
                    quantity = decimal_text(value)
            except InvalidOperation:
                pass
            if quantity is None:
                warning.append('數量須為正數且最多六位小數，請修正')
            items.append({'name': name or code or f'第 {number} 列', 'quantity': quantity,
                          'unit': text(row['unit']), 'product': product, 'processing': '',
                          'note': f'Excel 第 {number} 列', 'warning': '；'.join(warning)})
        if columns is None:
            raise Problem('請提供「品號」或「商品」欄，以及「數量」欄；可另加「單位」欄')
        if not items and not skipped:
            raise Problem('Excel 沒有可帶入的銷售明細')
    finally:
        book.close()
    fingerprint = hashlib.sha256(data).hexdigest()
    orders = [{
        'id': 'excel:' + fingerprint, 'sent_on': '', 'sent_at': '', 'customer': '',
        'location': filename, 'action': '新訂單', 'needs_review': False, 'source': [], 'items': items,
    }] if items else []
    return {'source': filename, 'standalone': True, 'skipped': skipped, 'orders': orders}
