import hashlib
import json
import warnings
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from zipfile import BadZipFile

import openpyxl

from .common import Problem, decimal_text, now, number, record, required_text, transaction
from .media import extract_media


def date_text(value):
    if value is None or value == '':
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    try:
        return date.fromisoformat(str(value).split(' ')[0]).isoformat()
    except ValueError:
        return None


def preview_import(paths, staging_dir):
    stage = Path(staging_dir)
    stage.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(b'inventory-import-v4')
    for path in sorted(paths, key=lambda p: p.name):
        digest.update(path.name.encode('utf-8'))
        with path.open('rb') as f:
            for chunk in iter(lambda: f.read(1024*1024), b''):
                digest.update(chunk)
    fingerprint = digest.hexdigest()
    cached = stage/(fingerprint+'.json')
    if cached.exists():
        return json.loads(cached.read_text('utf-8'))
    preview = {'import_id': fingerprint, 'fingerprint': fingerprint,
               'rows': [], 'issues': [], 'price_cards': [], 'invoice_items': [], 'files': [p.name for p in paths]}

    def issue(kind, source, message, original_value=None):
        preview['issues'].append({'kind': kind, 'source': source, 'message': message,
                                  'original_value': None if original_value is None else str(original_value)})

    def numeric(value, source, label):
        try:
            if value is not None and label in {'進價','售價'}:
                value = Decimal(str(value)).quantize(Decimal('0.000001'), rounding=ROUND_HALF_UP)
            return decimal_text(number(value, label, optional=True))
        except (Problem, InvalidOperation):
            issue('invalid_number', source, label+'不是有效數字，先保留為未設定', value)
            return None

    for path in paths:
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            if '庫存管理' in wb.sheetnames:
                rows = wb['庫存管理'].iter_rows(min_row=2, values_only=True)
                frows = formulas['庫存管理'].iter_rows(min_row=2, values_only=True)
                for i, (row, frow) in enumerate(zip(rows, frows), 2):
                    if len(row) < 13 or not row[3]:
                        continue
                    src = f'{path.name}/庫存管理/{i}'
                    for col in [5, 6, 11, 12]:
                        if row[col] is None and isinstance(frow[col], str) and frow[col].startswith('='):
                            issue('formula_cache_missing', src, f'第 {col+1} 欄公式缺少計算結果', frow[col])
                    item = {'source': src, 'category': str(row[0] or ''), 'supplier': str(row[1] or ''),
                            'code': str(row[2] or ''), 'name': str(row[3]), 'unit': str(row[4] or ''),
                            'cost': numeric(row[5], src, '進價'), 'price': numeric(row[6], src, '售價'),
                            'received_on': date_text(row[7]), 'expires_on': date_text(row[8]),
                            'quantity': numeric(row[11], src, '庫存'), 'minimum': numeric(row[12], src, '最低庫存')}
                    for idx, label in [(7, '進貨日期'), (8, '有效期限')]:
                        if row[idx] is not None and date_text(row[idx]) is None:
                            issue('invalid_date', src, label+'無法辨識，保留原始資料待確認', row[idx])
                    if item['quantity'] is None:
                        issue('unknown_stock', src, '庫存未盤點')
                    if item['price'] is None:
                        issue('missing_price', src, '售價尚未設定')
                    preview['rows'].append(item)
            elif '易發票系統' in wb.sheetnames:
                for i, row in enumerate(wb['易發票系統'].iter_rows(min_row=2, values_only=True), 2):
                    if not row[1]:
                        continue
                    preview['invoice_items'].append({'source': f'{path.name}/易發票系統/{i}',
                        'category': str(row[0]), 'code': str(row[1]), 'name': str(row[2]),
                        'price': numeric(row[3], str(i), '售價'), 'unit': str(row[4]),
                        'tax': row[5], 'kind': row[6]})
            else:
                images = extract_media(path, stage/'media')
                for ws in wb:
                    grid = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
                    cols = [6, 10] if ws.title == '1乳飲品' else [5, 11]
                    for r in (11, 17, 23, 29, 35, 41):
                        for col in cols:
                            name = grid.get((r-2, col))
                            if col == 10 and not name:
                                name = grid.get((r-2, 6))
                            price = grid.get((r, col))
                            if name is None and price is None:
                                continue
                            possible = [im for im in images if im['sheet'] == ws.title
                                        # The right photo may begin in column G and extend
                                        # across H:J; G belongs to the right-hand card.
                                        and r-4 <= im['row'] <= r and (im['col'] <= 6) == (col <= 7)]
                            preview['price_cards'].append({'source': f'{ws.title}!{openpyxl.utils.get_column_letter(col)}{r}',
                                'brand': str(grid.get((r-3, col)) or ''), 'name': str(name or ''),
                                'specification': str(grid.get((r-1, col)) or ''),
                                'original_price': str(price) if price is not None else None,
                                # Excel drawings are ordered back to front; an old image
                                # can remain behind its replacement in the same cell.
                                'image': possible[-1]['thumbnail_path'] if possible else None})
        finally:
            wb.close()
            formulas.close()
    by_code = defaultdict(list)
    for row in preview['rows']:
        by_code[row['code']].append(row)
    for code, rows in by_code.items():
        identities = {(r['name'], r['unit']) for r in rows}
        if len(identities) > 1 or not code:
            for row in rows:
                issue('code_collision', row['source'], f'品號 {code or "空白"} 對應多種商品，請確認品號')
        if len({r['price'] for r in rows if r['price'] is not None}) > 1:
            for row in rows:
                issue('price_conflict', row['source'], '同品號售價不一致，請確認')
    cached.write_text(json.dumps(preview, ensure_ascii=False), 'utf-8')
    return preview


def commit_import(conn, preview, resolutions, actor_id):
    fingerprint = preview['fingerprint']
    with transaction(conn):
        old = conn.execute('SELECT * FROM imports WHERE fingerprint=?', (fingerprint,)).fetchone()
        if old:
            return json.loads(old['preview'])['result']
        non_conflicts = [i for i in preview['issues'] if i['kind'] not in {'code_collision','price_conflict'}]
        if non_conflicts and resolutions.get('acknowledge_issues') is not True:
            raise Problem('請先閱讀待確認資料，並確認空值保留為未知；有誤的來源資料可先修正後再匯入')
        if conn.execute('SELECT 1 FROM products').fetchone():
            raise Problem('首次匯入只接受空資料庫，避免重複建立期初庫存', status=409)
        grouped = defaultdict(list)
        for source in preview['rows']:
            row = {**source}
            row['code'] = required_text(resolutions.get('codes', {}).get(row['source'], row['code']), '品號', 60)
            if row['source'] in resolutions.get('prices', {}):
                row['price'] = decimal_text(number(resolutions['prices'][row['source']], '售價', optional=True))
            grouped[row['code']].append(row)
        ids = {}
        for code, rows in grouped.items():
            if len({(r['name'], r['unit']) for r in rows}) > 1:
                raise Problem(f'品號 {code} 仍對應不同商品，請先修正', status=409)
            prices = {r['price'] for r in rows if r['price'] is not None}
            if len(prices) > 1:
                raise Problem(f'品號 {code} 的售價不一致，請先確認', status=409)
            first = rows[0]
            unit = required_text(first['unit'], '商品單位', 20)
            price = next(iter(prices), None)
            cur = conn.execute('''INSERT INTO products(code,name,category,supplier,unit,price,minimum,created_at,updated_at)
                                  VALUES(?,?,?,?,?,?,?,?,?)''',
                               (code, first['name'], first['category'], first['supplier'], unit, price, first['minimum'], now(), now()))
            pid = ids[code] = cur.lastrowid
            for r in rows:
                cur = conn.execute('''INSERT INTO batches(product_id,quantity,cost,received_on,expires_on,saleable,source)
                                      VALUES(?,?,?,?,?,?,?)''',
                                   (pid, r['quantity'], r['cost'], r['received_on'], r['expires_on'], int(bool(r['expires_on'])), r['source']))
                record(conn, pid, cur.lastrowid, 'opening', '原始 Excel 期初匯入', None, r['quantity'], actor_id)
        for card in preview['price_cards']:
            conn.execute('''INSERT INTO catalog_cards(source,brand,name,specification,original_price,image)
                            VALUES(:source,:brand,:name,:specification,:original_price,:image)''', card)
        for item in preview['invoice_items']:
            conn.execute('''INSERT INTO invoice_items(source,category,code,name,price,unit,tax,kind)
                            VALUES(:source,:category,:code,:name,:price,:unit,:tax,:kind)''', item)
        result = {'products': len(ids), 'batches': len(preview['rows']),
                  'catalog_pending': len(preview['price_cards']), 'invoice_pending': len(preview['invoice_items'])}
        conn.execute('INSERT INTO imports VALUES(?,?,?,?,?)',
                     (fingerprint, fingerprint, 'committed', json.dumps({'result': result}, ensure_ascii=False), now()))
        return result
