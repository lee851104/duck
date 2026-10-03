"""LINE 訂單帶入每日銷售單。

只有「已完成」、還沒扣過庫存的 LINE 訂單可以帶入。老闆在每日銷售單逐筆核對（選店內商品、確認數量），
帶入的訂單編號跟著銷售單一起送出：同一個交易裡標記，所以同一筆訂單不會扣兩次庫存；銷售單沖銷後可以重新帶入。
來源可以是系統裡的 LINE 接單，或 LINE 接單匯出、老闆在 Excel 改過數量／店內品號的檔案。
"""
import io
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from .common import Problem, decimal_text
from .line_orders import attach_details, rows_by_ids, sale_links

MAX_ORDERS = 300
MAX_EXCEL_BYTES = 10 * 1024 * 1024
MAX_EXCEL_ROWS = 5000
NO_ITEMS = '（沒有辨識出品項）'
STATUS_NAMES = {'open': '待處理', 'dismissed': '不是訂單'}


def product_view(p):
    return {k: p[k] for k in ('id', 'code', 'name', 'unit')} if p else None


def order_view(o, items):
    view = {k: o[k] for k in ('id', 'sent_on', 'sent_at', 'customer', 'location', 'action', 'carrier', 'note',
                               'review_reason')}
    view['needs_review'] = bool(o['needs_review'])
    view['source'] = [m['raw'] for m in o.get('source', [])]
    # 取消單不扣庫存，但仍跟著銷售單標記，之後不會再出現在待帶入清單。
    view['items'] = [] if o['action'] == '取消' else items
    return view


def pending_orders(conn):
    """系統裡已完成、還沒扣庫存的 LINE 訂單，依留言時間排列。"""
    rows = [dict(r) for r in conn.execute("SELECT * FROM line_orders WHERE status='done' ORDER BY sent_on,sent_at,id")]
    linked = sale_links(conn, [r['id'] for r in rows])
    rows = attach_details(conn, [r for r in rows if r['id'] not in linked])
    return [order_view(o, [{'name': i['name'], 'quantity': i['quantity'], 'unit': i['unit'],
                            'processing': i['processing'], 'note': i['note'],
                            'product': product_view(i['product'])} for i in o['items']]) for o in rows]


def quantity_text(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        return None
    return decimal_text(number) if number.is_finite() and number > 0 else None


def text(value):
    return '' if value is None else str(value).strip()


def read_rows(stream):
    """讀「訂單明細」工作表，回傳 [(列號, {欄名: 值})]；欄位用標題找，順序被調過也沒關係。"""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    data = stream.read(MAX_EXCEL_BYTES + 1)
    if len(data) > MAX_EXCEL_BYTES:
        raise Problem('檔案太大，請用 LINE 接單匯出的 Excel')
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, ValueError, OSError):
        raise Problem('讀不到這個 Excel 檔，請用 LINE 接單匯出的 .xlsx 檔') from None
    try:
        if '訂單明細' not in book.sheetnames:
            raise Problem('找不到「訂單明細」工作表，請用 LINE 接單匯出的 Excel')
        header, rows = None, []
        for number, values in enumerate(book['訂單明細'].iter_rows(values_only=True), 1):
            if number > MAX_EXCEL_ROWS:
                raise Problem(f'資料超過 {MAX_EXCEL_ROWS} 列，請縮小匯出的時間範圍')
            if header is None:
                names = [text(v) for v in values]
                if '訂單編號' in names and '商品' in names:
                    header = names
                elif number >= 10:
                    break
                continue
            row = {name: values[i] for i, name in enumerate(header) if name and i < len(values)}
            if any(v not in (None, '') for v in row.values()):
                rows.append((number, row))
        if header is None:
            raise Problem('這個 Excel 沒有「訂單編號」欄，請到 LINE 接單重新匯出一次')
        return rows
    finally:
        book.close()


def orders_from_excel(conn, stream):
    """Excel 的品項、數量、店內品號為準；訂單的地點、客人、原文和狀態以系統為準。"""
    products = {r['code']: dict(r) for r in conn.execute('SELECT id,code,name,unit FROM products')}
    grouped, skipped = {}, []
    for number, row in read_rows(stream):
        raw_id = text(row.get('訂單編號'))
        if not raw_id.isdigit():
            skipped.append({'id': None, 'kind': 'row', 'label': f'第 {number} 列', 'reason': '訂單編號空白或看不懂，略過這一列'})
            continue
        items = grouped.setdefault(int(raw_id), [])
        name = text(row.get('商品'))
        if not name or name == NO_ITEMS:
            continue
        code = text(row.get('店內品號'))
        product = products.get(code)
        items.append({'name': name, 'quantity': quantity_text(row.get('數量')), 'unit': text(row.get('單位')),
                      'processing': text(row.get('處理方式')), 'note': text(row.get('品項備註')),
                      'product': product_view(product),
                      'warning': f'店內品號「{code}」找不到，請重新選商品' if code and not product else ''})
    if len(grouped) > MAX_ORDERS:
        raise Problem(f'一次最多帶入 {MAX_ORDERS} 筆訂單，請縮小匯出的時間範圍')
    ids = list(grouped)
    found = {r['id']: dict(r) for r in rows_by_ids(conn, 'SELECT * FROM line_orders WHERE id IN ({})', ids)}
    linked = sale_links(conn, ids)
    ready = []
    for oid in ids:
        o = found.get(oid)
        label = f"{o['sent_on'][5:].replace('-', '/')} {o['sent_at']} {o['location'] or o['customer']}" if o else f'訂單編號 {oid}'
        if not o:
            skipped.append({'id': oid, 'kind': 'missing', 'label': label, 'reason': '系統裡找不到這筆訂單（可能是別台電腦的資料）'})
        elif o['status'] != 'done':
            skipped.append({'id': oid, 'kind': o['status'], 'label': label,
                            'reason': f"在 LINE 接單是「{STATUS_NAMES[o['status']]}」，按「已完成」後才能帶入"})
        elif oid in linked:
            skipped.append({'id': oid, 'kind': 'linked', 'label': label, 'reason': f'已帶入每日銷售單 #{linked[oid]}，不會重複扣庫存'})
        else:
            ready.append(o)
    attach_details(conn, ready)
    ready.sort(key=lambda o: (o['sent_on'], o['sent_at'], o['id']))
    return [order_view(o, grouped[o['id']]) for o in ready], skipped


def check_for_sheet(conn, ids):
    """送出銷售單前再確認一次：訂單都還是「已完成」，也還沒被別張銷售單帶入。"""
    if ids is None:
        return []
    if (not isinstance(ids, list) or len(ids) > MAX_ORDERS or any(type(i) is not int for i in ids)
            or len(set(ids)) != len(ids)):
        raise Problem('LINE 訂單格式不正確')
    found = {r['id']: r for r in rows_by_ids(conn, 'SELECT id,status,location,customer FROM line_orders WHERE id IN ({})', ids)}
    linked = sale_links(conn, ids)
    for oid in ids:
        o = found.get(oid)
        if not o:
            raise Problem('找不到部分 LINE 訂單，請取消帶入後重新帶入', status=409)
        name = o['location'] or o['customer']
        if o['status'] != 'done':
            raise Problem(f'LINE 訂單「{name}」已不是「已完成」，請取消帶入後重新帶入', status=409)
        if oid in linked:
            raise Problem(f'LINE 訂單「{name}」已帶入每日銷售單 #{linked[oid]}，請取消帶入後重新帶入', status=409)
    return sorted(ids)
