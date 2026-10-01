import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from app.common import Problem
from app.db import connect_db, init_db
from app.imports import preview_import, commit_import


class ImportTest(unittest.TestCase):
    def test_catalog_uses_frontmost_image_when_excel_overlays_old_photo(self):
        wb=Workbook();ws=wb.active;ws.title='11調味粉'
        ws['E8']='台糖';ws['E9']='晶冰糖';ws['E10']='1kg／包';ws['E11']=50
        path=self.root/'overlaid.xlsx';wb.save(path)
        images=[{'sheet':ws.title,'row':8,'col':2,'thumbnail_path':'old-sugar.jpg'},
                {'sheet':ws.title,'row':8,'col':2,'thumbnail_path':'front-rock-sugar.jpg'}]
        with patch('app.imports.extract_media',return_value=images):
            p=preview_import([path],self.root/'overlay-stage')
        self.assertEqual(p['price_cards'][0]['image'],'front-rock-sugar.jpg')

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.conn = connect_db(self.root/'db.sqlite')
        self.addCleanup(self.conn.close)
        init_db(self.conn)
        wb = Workbook()
        ws = wb.active
        ws.title = '庫存管理'
        ws.append(['類別', '供應商', '品號', '品名', '單位', '進價', '售價', '進貨日期', '有效期限', '剩餘', '管理', '庫存', '最低庫存'])
        ws.append(['米', '供應商', 'I1', '芋香米', '包', 100, 140, None, None, None, None, None, 2])
        ws.append(['米', '供應商', 'I1', '九號米', '包', 100, None, None, None, None, None, 2, 2])
        ws.append(['肉', '供應商', 'D1', '貢丸', '包', 100, 150, None, '2027-01-01', None, None, 2, 3])
        ws.append(['肉', '供應商', 'D1', '貢丸', '包', '=1+2', 150, None, '2027-02-01', None, None, 7, 3])
        self.path = self.root/'inventory.xlsx'
        wb.save(self.path)

    def test_preview_preserves_unknown_and_collision(self):
        p = preview_import([self.path], self.root/'stage')
        self.assertEqual(len(p['rows']), 4)
        self.assertIsNone(p['rows'][0]['quantity'])
        self.assertIsNone(p['rows'][1]['price'])
        self.assertTrue(any(i['kind'] == 'code_collision' for i in p['issues']))
        self.assertTrue(any(i['kind'] == 'formula_cache_missing' for i in p['issues']))
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM products').fetchone()[0], 0)

    def test_commit_requires_collision_resolution(self):
        p = preview_import([self.path], self.root/'stage')
        with self.assertRaises(Problem):
            commit_import(self.conn, p, {'acknowledge_issues':True}, 1)

    def test_commit_keeps_batches_and_is_idempotent(self):
        p = preview_import([self.path], self.root/'stage')
        resolutions = {'codes': {p['rows'][1]['source']: 'I2'}, 'acknowledge_issues': True}
        a = commit_import(self.conn, p, resolutions, 1)
        b = commit_import(self.conn, p, resolutions, 1)
        self.assertEqual(a, b)
        self.assertEqual(a['products'], 3)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM batches').fetchone()[0], 4)
        self.assertIsNone(self.conn.execute('SELECT quantity FROM batches ORDER BY id').fetchone()[0])

    def test_collision_resolved_to_existing_other_product_rejected(self):
        p = preview_import([self.path], self.root/'stage')
        with self.assertRaises(Problem):
            commit_import(self.conn, p, {'codes': {p['rows'][1]['source']: 'D1'},'acknowledge_issues':True}, 1)

    def test_fractional_cost_from_excel_is_preserved(self):
        wb = Workbook()
        ws = wb.active
        ws.title = '庫存管理'
        ws.append(['類別','供應商','品號','品名','單位','進價','售價','進貨日期','效期','天數','管理','庫存','最低'])
        ws.append(['調味','供應商','P3','胡椒','瓶',33.3333333333333,40,None,None,None,None,1,2])
        file = self.root/'fractions.xlsx'
        wb.save(file)
        p = preview_import([file], self.root/'stage')
        self.assertEqual(p['rows'][0]['cost'], '33.333333')

    def test_unknown_values_require_explicit_acknowledgement(self):
        p = preview_import([self.path], self.root/'stage')
        with self.assertRaises(Problem):
            commit_import(self.conn, p, {'codes':{p['rows'][1]['source']:'I2'}},1)

    def test_invalid_date_keeps_original_value_in_issue(self):
        wb = Workbook()
        ws = wb.active
        ws.title='庫存管理'
        ws.append(['類別','供應商','品號','品名','單位','進價','售價','進貨日期','效期','天數','管理','庫存','最低'])
        ws.append(['食品','供應商','BADDATE','測試','包',10,20,None,'不是日期',None,None,3,1])
        file=self.root/'bad-date.xlsx';wb.save(file)
        p=preview_import([file],self.root/'stage')
        issue=next(i for i in p['issues'] if i['kind']=='invalid_date')
        self.assertEqual(issue['original_value'],'不是日期')
