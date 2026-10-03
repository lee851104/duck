import io
import json
import os
import unittest
from unittest import mock
from uuid import uuid4

from app import line_chat as lc
from app.db import connect_db
from app.products import create_product
from tests import test_app

BASE = """[LINE] 聊天記錄
2026.10.01 星期四
10:44 菜騎鴨官方小幫手0426567066  (emoji)今日菜單(2026.10.01)(emoji)
(emoji)葉菜類
   蔥 35/把
11:00 Test Alice 測試社區甲

蔥 35/把*1 切
台灣豬五花*1
11:02 小幫手 圖片
11:06 Test Bob 可以看看葡萄的照片嗎
"""

APPEND = """2026.10.02 星期五
09:10 Test Chen 測試社區乙12號
黃玉米 10/2支(剁、去皮)
/TEST123
09:11 Test Lin。（測試社區丙6D3) 圖片
09:12 Test Lin。（測試社區丙6D3)已收回訊息
"""


def order(source_ids, items, **values):
    return {'source_ids': source_ids, 'customer': values.get('customer', '客人'),
            'location': values.get('location', ''), 'action': values.get('action', '新訂單'),
            'items': items, 'carrier': values.get('carrier', ''), 'payment': '', 'note': '',
            'needs_review': values.get('needs_review', False), 'review_reason': ''}


def item(name, quantity, product_id=None, processing=''):
    return {'product_id': product_id, 'name': name, 'quantity': quantity, 'unit': '包',
            'unit_price': None, 'processing': processing, 'note': ''}


class ParseTest(unittest.TestCase):
    def tail(self, text):
        return [m.fp for m in lc.parse_chat(text)][-lc.TAIL_SIZE:]

    def test_messages_keep_multiline_text_and_skip_header(self):
        messages = lc.parse_chat(BASE)
        self.assertEqual([m.time for m in messages], ['10:44', '11:00', '11:02', '11:06'])
        self.assertEqual(messages[1].raw, '11:00 Test Alice 測試社區甲\n\n蔥 35/把*1 切\n台灣豬五花*1')
        self.assertEqual(lc.parse_chat('2026/10/01（四）\n下午1:05\t王\t蔥+1\n')[0].time, '13:05')
        self.assertEqual(lc.decode_export(b'\xef\xbb\xbf' + BASE.encode()).splitlines()[1], '2026.10.01 星期四')

    def test_only_messages_after_last_position_are_new(self):
        new, missing = lc.find_new_messages(self.tail(BASE), lc.parse_chat(BASE + APPEND))
        self.assertEqual(([m.time for m in new], missing), (['09:10', '09:11', '09:12'], 0))
        repeat = BASE + '11:06 Test Bob 可以看看葡萄的照片嗎\n'
        self.assertEqual(len(lc.find_new_messages(self.tail(BASE), lc.parse_chat(repeat))[0]), 1)

    def test_recalled_last_message_and_unrelated_chat(self):
        old = BASE + '11:09 Test Wu 西瓜+1\n'
        new, missing = lc.find_new_messages(self.tail(old), lc.parse_chat(BASE + '11:10 Test Wu已收回訊息\n'))
        self.assertEqual(([m.time for m in new], missing), (['11:10'], 1))
        other = lc.parse_chat('2026.10.03 星期六\n10:00 X 你好\n10:01 Y 嗨\n')
        self.assertIsNone(lc.find_new_messages(self.tail(BASE), other))

    def test_classify_menu_and_chunks(self):
        kinds = [lc.classify(m) for m in lc.parse_chat(BASE + APPEND)]
        self.assertEqual(kinds, ['staff', 'customer', 'staff', 'customer', 'customer', 'image', 'recall'])
        menu = lc.latest_menu(lc.parse_chat(BASE))
        self.assertIn('蔥 35/把', menu['text'])
        self.assertNotIn('(emoji)', menu['text'])
        chunks = lc.chunk_blocks([(i, 'x' * 40) for i in range(1, 6)], 100)
        self.assertEqual([[i for i, _ in c] for c in chunks], [[1, 2], [3, 4], [5]])

    def test_cost_estimate(self):
        usage = lc.add_usage({}, {'input_tokens': 1000, 'output_tokens': 200,
                                  'input_tokens_details': {'cached_tokens': 400}})
        self.assertAlmostEqual(lc.estimate_cost('gpt-6-luna', usage), (600 * 0.1 + 400 * 0.01 + 200 * 0.5) / 1e6)
        self.assertIsNone(lc.estimate_cost('unknown-model', usage))


class LineOrdersApiTest(unittest.TestCase):
    login = test_app.AppTest.login

    def setUp(self):
        environment = mock.patch.dict(os.environ, {}, clear=False)
        environment.start()
        self.addCleanup(environment.stop)
        for name in ('OPENAI_API_KEY', 'OPENAI_MODEL', 'OPENAI_REASONING_EFFORT'):
            os.environ.pop(name, None)
        test_app.AppTest.setUp(self)
        self.app.config['ENABLE_LINE_ORDERS'] = True
        self.sent, self.reply, self.failure = [], (lambda payload: []), None
        self.app.config['OPENAI_TRANSPORT'] = self.fake
        self.login()
        self.headers = {'X-CSRF-Token': self.token}
        self.conn = connect_db(self.app.config['DB_PATH'])
        self.addCleanup(self.conn.close)
        self.pork = create_product(self.conn, {'code': 'B01', 'name': '台灣豬五花(火)', 'unit': '包',
                                               'price': '200'}, 1)['id']

    def fake(self, payload):
        self.sent.append(payload)
        if self.failure:
            raise self.failure
        text = json.dumps({'orders': self.reply(payload)}, ensure_ascii=False)
        return {'status': 'completed', 'usage': {'input_tokens': 1000, 'output_tokens': 200},
                'output': [{'type': 'reasoning'}, {'type': 'message', 'content': [{'type': 'output_text', 'text': text}]}]}

    def upload(self, text, start=None, name='[LINE]菜騎鴨.txt'):
        data = {'file': (io.BytesIO(text.encode('utf-8')), name)}
        if start is not None:
            data['start'] = start
        return self.client.post('/api/line-orders/import', data=data, headers=self.headers,
                                content_type='multipart/form-data')

    def listing(self, status='open'):
        response = self.client.get('/api/line-orders?status=' + status)
        self.assertEqual(response.status_code, 200, response.get_json())
        return response.get_json()

    def test_first_import_asks_where_to_start(self):
        response = self.upload(BASE)
        self.assertEqual(response.status_code, 409)
        error = response.get_json()['error']
        self.assertEqual(error['code'], 'line_first_import')
        self.assertEqual((error['fields']['last_on'], error['fields']['last_at']), ('2026-10-01', '11:06'))
        self.assertEqual(self.sent, [])

    def test_baseline_then_only_new_messages_become_orders(self):
        first = self.upload(BASE, 'baseline').get_json()
        self.assertEqual((first['status'], first['new_messages'], self.sent), ('baseline', 0, []))
        self.reply = lambda payload: [order([1], [item('黃玉米', 2, processing='剁、去皮')], location='測試社區乙12號',
                                            carrier='/TEST123')]
        second = self.upload(BASE + APPEND)
        self.assertEqual(second.status_code, 200, second.get_json())
        result = second.get_json()
        self.assertEqual((result['new_messages'], result['customer_messages'], result['orders_created']), (3, 1, 1))
        self.assertEqual((len(result['import']['images']), len(result['import']['recalls'])), (1, 1))
        system, user = (m['content'] for m in self.sent[0]['input'])
        self.assertIn(f'#{self.pork} 台灣豬五花(火)｜包｜$200', system)
        self.assertIn('蔥 35/把', user)
        self.assertIn('[1] 2026-10-02 09:10 Test Chen 測試社區乙12號', user)
        self.assertNotIn('Alice', user)
        self.assertEqual(self.sent[0]['text']['format']['type'], 'json_schema')
        self.assertFalse(self.sent[0]['store'])

        again = self.upload(BASE + APPEND).get_json()
        self.assertEqual((again['import_id'], again['new_messages'], len(self.sent)), (None, 0, 1))
        data = self.listing()
        self.assertEqual(data['counts'], {'open': 1, 'done': 0, 'dismissed': 0})
        placed = data['items'][0]
        self.assertEqual((placed['sent_on'], placed['sent_at'], placed['carrier']), ('2026-10-02', '09:10', '/TEST123'))
        self.assertEqual(placed['source'][0]['raw'].splitlines()[1], '黃玉米 10/2支(剁、去皮)')
        self.assertEqual(data['summary'][0]['name'], '黃玉米')
        self.assertEqual(data['summary'][0]['total'], '2')
        self.assertEqual(data['latest']['orders'], 1)
        self.assertTrue(placed['is_new'])
        self.assertEqual(data['new_count'], 1)
        # 畫面只顯示有沒有接上，不給模型名稱或費用
        self.assertEqual(set(data['ai']), {'configured', 'cloud', 'key_location'})
        self.assertIsNotNone(self.conn.execute('SELECT cost FROM line_imports WHERE status="done"').fetchone()[0])

    def test_open_orders_sorted_by_time_both_ways(self):
        self.upload(BASE, 'baseline')
        self.reply = lambda payload: [order([1], [item('黃玉米', 2)], location='第一批')]
        self.upload(BASE + APPEND)
        later = '2026.10.03 星期六\n08:00 王小美 測試社區丁B10-7\n蔥+1\n'
        self.reply = lambda payload: [order([1], [item('蔥', 1)], location='第二批')]
        self.upload(BASE + APPEND + later)
        data = self.listing()  # 預設新到舊
        self.assertEqual([(o['location'], o['is_new']) for o in data['items']], [('第二批', True), ('第一批', False)])
        self.assertEqual(data['new_count'], 1)
        oldest = self.client.get('/api/line-orders?status=open&sort=old').get_json()
        self.assertEqual([o['location'] for o in oldest['items']], ['第一批', '第二批'])
        self.assertEqual(self.client.get('/api/line-orders?sort=random').status_code, 400)

    def test_time_range_filters_list_counts_and_totals(self):
        self.upload(BASE, 'baseline')
        self.reply = lambda payload: [order([1], [item('黃玉米', 2)], location='10/02 的單')]
        self.upload(BASE + APPEND)
        later = '2026.10.03 星期六\n08:00 王小美 測試社區丁B10-7\n黃玉米+3\n'
        self.reply = lambda payload: [order([1], [item('黃玉米', 3)], location='10/03 的單')]
        self.upload(BASE + APPEND + later)
        day = lambda d, status='open': self.client.get(
            f'/api/line-orders?status={status}&from={d}T00:00&to={d}T23:59').get_json()
        today = day('2026-10-03')
        self.assertEqual([o['location'] for o in today['items']], ['10/03 的單'])
        self.assertEqual(today['counts'], {'open': 1, 'done': 0, 'dismissed': 0})
        self.assertEqual([(g['name'], g['total']) for g in today['summary']], [('黃玉米', '3')])
        everything = self.listing()
        self.assertEqual(everything['summary'][0]['total'], '5')
        self.assertIsNone(self.listing('done')['summary'])  # 全部時間的已完成不加總
        bad = self.client.get('/api/line-orders?from=2026-10-03T10:00&to=2026-10-03T09:00')
        self.assertEqual(bad.status_code, 400)

    def test_export_excel_for_a_day(self):
        from openpyxl import load_workbook
        self.upload(BASE, 'baseline')
        tricky = '2026.10.02 星期五\n09:10 =HYPERLINK("http://x") 測試社區丁B6-3\n黃玉米+2\n09:15 阿明 請問有西瓜嗎\n'
        self.reply = lambda payload: [order([1], [item('黃玉米', 2, processing='剁')], location='測試社區丁B6-3',
                                            customer='=HYPERLINK("http://x")', needs_review=True)]
        self.upload(BASE + tricky)
        response = self.client.get('/api/line-orders/export?from=2026-10-02T00:00&to=2026-10-02T23:59')
        self.assertEqual(response.status_code, 200)
        self.assertIn("LINE%E8%A8%82%E5%96%AE_2026-10-02.xlsx", response.headers['Content-Disposition'])
        book = load_workbook(io.BytesIO(response.data))
        self.assertEqual(book.sheetnames, ['商品合計', '訂單明細', '其他留言'])
        totals, detail, others = (book[name] for name in book.sheetnames)
        self.assertEqual(totals['A1'].value, 'LINE 訂單｜2026-10-02')
        self.assertEqual([c.value for c in totals[4]][:6], ['黃玉米', 2, '包', 1, None, 1])
        self.assertEqual(detail['D4'].value, '=HYPERLINK("http://x")')
        self.assertEqual(detail['D4'].data_type, 's')  # 不會變成公式
        self.assertIn('黃玉米+2', detail['Q4'].value)
        self.assertEqual([c.value for c in detail[3]][-2:], ['訂單編號', '已扣庫存銷售單'])
        self.assertEqual(detail['R4'].value, self.listing()['items'][0]['id'])
        self.assertEqual([(c.value) for c in others[4]][2:], ['AI 沒當成訂單，請看一下', '09:15 阿明 請問有西瓜嗎'])
        empty = self.client.get('/api/line-orders/export?from=2026-09-01T00:00&to=2026-09-01T23:59')
        self.assertEqual(empty.status_code, 404)

    def test_start_time_and_matching_existing_products(self):
        self.reply = lambda payload: [order([1], [item('台灣豬五花(火)', 1, self.pork), item('蔥', 1, 99999, '切')],
                                            location='測試社區甲')]
        result = self.upload(BASE, '2026-10-01T11:00').get_json()
        self.assertEqual((result['new_messages'], result['customer_messages'], result['orders_created']), (3, 2, 1))
        placed = self.listing()['items'][0]
        pork, onion = placed['items']
        self.assertEqual((pork['product']['code'], pork['product']['name']), ('B01', '台灣豬五花(火)'))
        self.assertIsNone(onion['product_id'])  # 編號不在店內商品清單內就不採用
        detail = self.client.get(f"/api/line-orders/imports/{result['import_id']}").get_json()
        self.assertEqual([m['sent_at'] for m in detail['not_orders']], ['11:06'])

    def test_renamed_export_continues_the_same_chat(self):
        self.upload(BASE, 'baseline')
        self.reply = lambda payload: [order([1], [item('黃玉米', 2)])]
        result = self.upload(BASE + APPEND, name='[LINE]菜騎鴨(更新1).txt')
        self.assertEqual(result.status_code, 200, result.get_json())
        self.assertEqual((result.get_json()['chat'], result.get_json()['new_messages']), ('菜騎鴨', 3))
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM line_chats').fetchone()[0], 1)

    def test_older_or_unrelated_exports_are_refused(self):
        self.upload(BASE + APPEND, 'baseline')
        self.assertEqual(self.upload(BASE).get_json()['error']['code'], 'line_older_export')
        unrelated = self.upload('2026.10.03 星期六\n10:00 X 你好\n10:01 Y 嗨\n')
        self.assertEqual(unrelated.get_json()['error']['code'], 'line_position_lost')
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM line_messages').fetchone()[0], 0)

    def test_ai_failure_keeps_messages_and_can_retry(self):
        self.upload(BASE, 'baseline')
        self.failure = lc.AIError('OpenAI 額度（credits）不足，請到 OpenAI 帳戶加值')
        result = self.upload(BASE + APPEND).get_json()
        self.assertIn('額度', result['convert_error'])
        self.assertEqual(result['import']['status'], 'failed')
        self.failure = None
        self.reply = lambda payload: [order([1], [item('黃玉米', 2)])]
        retry = self.client.post(f"/api/line-orders/imports/{result['import_id']}/convert", headers=self.headers)
        self.assertEqual(retry.status_code, 200, retry.get_json())
        self.assertEqual((retry.get_json()['orders_created'], retry.get_json()['import']['status']), (1, 'done'))
        again = self.client.post(f"/api/line-orders/imports/{result['import_id']}/convert", headers=self.headers)
        self.assertEqual(again.status_code, 409)

    def test_missing_key_leaves_import_pending(self):
        self.app.config['OPENAI_TRANSPORT'] = None
        self.upload(BASE, 'baseline')
        result = self.upload(BASE + APPEND).get_json()
        self.assertIn('金鑰', result['convert_error'])
        self.assertEqual(result['import']['status'], 'pending')
        self.assertFalse(self.listing()['ai']['configured'])
        (self.app.config['DATA_DIR'] / 'openai-key.txt').write_text('sk-test-key\n', encoding='utf-8')
        self.assertTrue(self.listing()['ai']['configured'])

    def test_restart_from_time_skips_messages_already_imported(self):
        self.upload(BASE, 'baseline')
        self.reply = lambda payload: [order([1], [item('黃玉米', 2)])]
        self.upload(BASE + APPEND)
        restart = self.upload(BASE + APPEND, '2026-10-02T00:00').get_json()
        self.assertEqual((restart['new_messages'], restart['status']), (0, 'baseline'))
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM line_orders').fetchone()[0], 1)

    def test_status_updates_use_versions(self):
        self.upload(BASE, 'baseline')
        self.reply = lambda payload: [order([1], [item('黃玉米', 2)])]
        self.upload(BASE + APPEND)
        oid = self.listing()['items'][0]['id']
        patch = lambda body: self.client.patch(f'/api/line-orders/{oid}', json=body, headers=self.headers)
        self.assertEqual(patch({'status': 'done', 'version': 1}).get_json()['version'], 2)
        self.assertEqual(patch({'status': 'open', 'version': 1}).status_code, 409)
        self.assertEqual(patch({'status': 'shipped', 'version': 2}).status_code, 400)
        self.assertEqual(self.listing('done')['total'], 1)
        self.assertEqual(self.listing()['summary'], [])

    def test_boss_can_edit_an_order(self):
        self.upload(BASE, 'baseline')
        self.reply = lambda payload: [order([1], [item('黃玉米', None)], location='', needs_review=True)]
        self.upload(BASE + APPEND)
        placed = self.listing()['items'][0]
        self.assertTrue(placed['needs_review'])
        self.assertIsNone(placed['edited_at'])
        products = self.client.get('/api/line-orders/products').get_json()['items']
        self.assertEqual([(p['id'], p['code']) for p in products], [(self.pork, 'B01')])
        edit = lambda body: self.client.put(f"/api/line-orders/{placed['id']}", json=body, headers=self.headers)
        body = {'version': placed['version'], 'location': ' 測試社區乙12號 ', 'customer': 'Test Chen', 'action': '新訂單',
                'carrier': '/TEST123', 'payment': '', 'note': '下午送', 'needs_review': False, 'review_reason': '還在',
                'items': [{'name': '黃玉米', 'quantity': '2', 'unit': '支', 'processing': '剁、去皮', 'note': '',
                           'product_id': None, 'unit_price': None},
                          {'name': '台灣豬五花(火)', 'quantity': '1', 'unit': '包', 'processing': '', 'note': '',
                           'product_id': self.pork, 'unit_price': '200'}]}
        response = edit(body)
        self.assertEqual(response.status_code, 200, response.get_json())
        saved = self.listing()['items'][0]
        self.assertEqual((saved['location'], saved['note'], saved['needs_review'], saved['review_reason']),
                         ('測試社區乙12號', '下午送', False, ''))
        self.assertIsNotNone(saved['edited_at'])
        self.assertEqual([(i['name'], i['quantity'], i['unit'], i['processing']) for i in saved['items']],
                         [('黃玉米', '2', '支', '剁、去皮'), ('台灣豬五花(火)', '1', '包', '')])
        self.assertEqual(saved['items'][1]['product']['code'], 'B01')
        totals = {g['name']: g['total'] for g in self.listing()['summary']}
        self.assertEqual(totals, {'黃玉米': '2', '台灣豬五花(火)': '1'})  # 合計跟著改
        self.assertEqual(edit(body).status_code, 409)  # 版本已更新
        body['version'] = saved['version']
        for change in ({'action': '退貨'}, {'items': [{'name': ' '}]}, {'items': [{'name': '蔥', 'product_id': 99999}]},
                       {'items': [{'name': '蔥', 'quantity': '-1'}]}, {'needs_review': 'yes'}, {'location': 'x' * 101}):
            self.assertEqual(edit({**body, **change}).status_code, 400, change)
        kept = edit({**body, 'items': [], 'needs_review': True, 'review_reason': '客人說再確認'})
        self.assertEqual(kept.status_code, 200, kept.get_json())
        self.assertEqual((kept.get_json()['items'], kept.get_json()['review_reason']), ([], '客人說再確認'))

    def test_old_database_gains_edited_column(self):
        import sqlite3
        from app.db import init_db
        from pathlib import Path
        schema = Path(test_app.__file__).parents[1].joinpath('app', 'schema.sql').read_text('utf-8')
        old = schema.replace(',\n edited_at TEXT);', ');')
        self.assertNotEqual(old, schema)
        conn = sqlite3.connect(':memory:')
        conn.executescript(old)
        init_db(conn)
        self.assertIn('edited_at', [r[1] for r in conn.execute('PRAGMA table_info(line_orders)')])
        conn.close()

    def test_login_and_csrf_are_required(self):
        anonymous = self.app.test_client()
        self.assertEqual(anonymous.get('/api/line-orders').status_code, 401)
        data = {'file': (io.BytesIO(BASE.encode()), '[LINE]x.txt'), 'start': 'baseline'}
        response = self.client.post('/api/line-orders/import', data=data, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 403)


TWO_ORDERS = BASE + """2026.10.02 星期五
09:10 甲 測試社區甲
台灣豬五花*2
蔥+1
09:20 乙 測試社區丁
黃玉米+1
"""


class LineToDailySalesTest(unittest.TestCase):
    """已完成的 LINE 訂單逐筆核對後帶入每日銷售單；同一筆訂單只會扣一次庫存。"""
    login = test_app.AppTest.login
    setUp, fake, upload, listing = (LineOrdersApiTest.setUp, LineOrdersApiTest.fake,
                                    LineOrdersApiTest.upload, LineOrdersApiTest.listing)

    def make_orders(self):
        self.conn.execute("UPDATE batches SET quantity='10',saleable=1,expires_on='2027-01-01' WHERE product_id=?",
                          (self.pork,))
        self.onion = create_product(self.conn, {'code': 'V01', 'name': '蔥', 'unit': '把'}, 1)['id']
        self.upload(BASE, 'baseline')
        self.reply = lambda payload: [
            order([1], [item('台灣豬五花(火)', 2, self.pork), item('蔥', 1)], location='測試社區甲'),
            order([2], [item('黃玉米', 1)], location='測試社區丁')]
        self.upload(TWO_ORDERS)
        return [o['id'] for o in self.client.get('/api/line-orders?sort=old').get_json()['items']]

    def mark(self, oid, status, expect=200):
        version = self.conn.execute('SELECT version FROM line_orders WHERE id=?', (oid,)).fetchone()[0]
        response = self.client.patch(f'/api/line-orders/{oid}', json={'status': status, 'version': version},
                                     headers=self.headers)
        self.assertEqual(response.status_code, expect, response.get_json())
        return response

    def pending(self):
        return self.client.get('/api/daily-sales/line-orders').get_json()['orders']

    def sheet(self, line_orders, sold_on='2026-10-01', quantity='2', expect=200):
        body = {'sold_on': sold_on, 'note': '', 'items': [{'product_id': self.pork, 'quantity': quantity}],
                'line_orders': line_orders}
        preview = self.client.post('/api/daily-sales/preview', json=body, headers=self.headers)
        if preview.status_code != 200:
            self.assertEqual(preview.status_code, expect, preview.get_json())
            return preview
        body.update(revision=preview.get_json()['revision'], request_id=str(uuid4()))
        response = self.client.post('/api/daily-sales', json=body, headers=self.headers)
        self.assertEqual(response.status_code, expect, response.get_json())
        return response

    def excel(self, data, name='LINE訂單.xlsx'):
        return self.client.post('/api/daily-sales/line-orders/excel', headers=self.headers,
                                data={'file': (io.BytesIO(data), name)}, content_type='multipart/form-data')

    def test_only_done_orders_once_and_void_releases_them(self):
        first, second = self.make_orders()
        self.assertEqual(self.pending(), [])  # 待處理的不能帶入
        self.mark(first, 'done')
        self.mark(second, 'done')
        orders = self.pending()
        self.assertEqual([o['location'] for o in orders], ['測試社區甲', '測試社區丁'])
        pork, onion = orders[0]['items']
        self.assertEqual((pork['product']['code'], pork['quantity'], onion['product']), ('B01', '2', None))
        self.assertIn('台灣豬五花*2', orders[0]['source'][0])

        sheet = self.sheet([first, second]).get_json()
        self.assertEqual([o['location'] for o in sheet['line_orders']], ['測試社區甲', '測試社區丁'])
        self.assertEqual(self.pending(), [])
        self.assertEqual({o['sale_id'] for o in self.listing('done')['items']}, {sheet['id']})
        self.assertIn('沖銷', self.mark(first, 'open', 409).get_json()['error']['message'])
        locked = self.client.put(f'/api/line-orders/{first}', headers=self.headers, json={
            'version': 99, 'action': '新訂單', 'needs_review': False, 'items': []})
        self.assertEqual(locked.status_code, 409)  # 已扣庫存的訂單不能改
        again = self.sheet([first], sold_on='2026-09-30', expect=409)
        self.assertIn(f"#{sheet['id']}", again.get_json()['error']['message'])

        void = self.client.post(f"/api/daily-sales/{sheet['id']}/void",
                                json={'reason': '數量填錯', 'request_id': str(uuid4())}, headers=self.headers)
        self.assertEqual(void.status_code, 200, void.get_json())
        self.assertEqual(len(self.pending()), 2)  # 沖銷後可以重新帶入
        self.mark(first, 'open')
        self.assertEqual(self.sheet([first], expect=409).status_code, 409)  # 已不是「已完成」
        self.assertEqual(self.sheet([second, second], expect=400).status_code, 400)

    def test_excel_uses_edited_quantities_and_codes(self):
        from openpyxl import load_workbook
        first, second = self.make_orders()
        self.mark(first, 'done')
        book = load_workbook(io.BytesIO(self.client.get('/api/line-orders/export').data))
        detail = book['訂單明細']
        column = {c.value: c.column for c in detail[3]}
        for row in detail.iter_rows(min_row=4):
            name = row[column['商品'] - 1].value
            if name == '台灣豬五花(火)':
                row[column['數量'] - 1].value = 3  # 老闆在 Excel 改數量
            if name == '蔥':
                row[column['店內品號'] - 1].value = 'V01'  # 補上店內品號
            if name == '黃玉米':
                row[column['店內品號'] - 1].value = 'NOPE'
        detail.append(['2026-10-02'] + [None] * (column['訂單編號'] - 2) + ['abc'])
        buffer = io.BytesIO()
        book.save(buffer)
        data = buffer.getvalue()

        result = self.excel(data).get_json()
        self.assertEqual([o['id'] for o in result['orders']], [first])
        pork, onion = result['orders'][0]['items']
        self.assertEqual((pork['quantity'], pork['product']['code']), ('3', 'B01'))
        self.assertEqual((onion['product']['name'], onion['product']['unit']), ('蔥', '把'))
        reasons = [s['reason'] for s in result['skipped']]
        self.assertTrue(any('看不懂' in r for r in reasons))
        self.assertTrue(any('待處理' in r for r in reasons))  # 第二筆還沒按已完成

        self.sheet([first], quantity='3')
        skipped = self.excel(data).get_json()
        self.assertEqual(skipped['orders'], [])
        self.assertTrue(any('不會重複扣庫存' in s['reason'] for s in skipped['skipped']))
        self.assertEqual(self.conn.execute('SELECT quantity FROM batches WHERE product_id=?', (self.pork,)).fetchone()[0], '7')

    def test_excel_rejects_other_files(self):
        from openpyxl import Workbook
        self.assertEqual(self.excel(b'not excel').status_code, 400)
        self.assertEqual(self.excel(b'x', 'orders.csv').status_code, 400)
        book = Workbook()
        book.active.title = '訂單明細'
        book.active.append(['日期', '商品', '數量'])
        buffer = io.BytesIO()
        book.save(buffer)
        response = self.excel(buffer.getvalue())
        self.assertEqual(response.status_code, 400)
        self.assertIn('訂單編號', response.get_json()['error']['message'])
        anonymous = self.app.test_client()
        self.assertEqual(anonymous.get('/api/daily-sales/line-orders').status_code, 401)

if __name__ == '__main__':
    unittest.main()
