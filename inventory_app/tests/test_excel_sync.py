import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from app import create_app
from app.common import Problem
from app.db import connect_db
from app.excel_sync import FILENAMES, write_catalog
from app.products import create_product


class ExcelSyncTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        source = self.root/'raw'
        source.mkdir()
        self.template = source/'商品價目表.xlsx'
        workbook = Workbook()
        workbook.active.title = '價目'
        workbook.active['E11'] = 20
        workbook.active['E9'] = '測試'
        workbook.active['K11'] = 80
        workbook.save(self.template)
        self.app = create_app({'TESTING': True, 'DATA_DIR': self.root/'data',
                               'SOURCE_DIR': source, 'SECRET_KEY': 'test', 'TODAY': '2026-10-02'})
        self.sync = self.app.extensions['excel_sync']
        self.conn = connect_db(self.app.config['DB_PATH'])
        self.addCleanup(self.conn.close)
        self.product = create_product(self.conn, {'code': '0001', 'name': '=商品', 'unit': '包', 'price': '45'}, 1)
        self.conn.execute('''INSERT INTO catalog_cards(source,name,original_price,product_id,pricing_mode)
                             VALUES('價目!E11','測試','20',?,'product')''', (self.product['id'],))
        self.conn.execute("INSERT INTO catalog_cards(source,name,original_price,independent) VALUES('價目!K11','整箱','80',1)")
        self.conn.execute('''INSERT INTO invoice_items(source,category,code,name,price,unit,tax,kind,product_id)
                            VALUES('invoice','1734','000099','舊名','20','包',1,0,?)''', (self.product['id'],))

    def rows(self, folder, kind):
        workbook = load_workbook(Path(folder)/FILENAMES[kind], data_only=False)
        self.addCleanup(workbook.close)
        return list(workbook.active.values)

    def test_snapshot_preserves_unknown_zero_batches_and_identifier(self):
        self.conn.execute("INSERT INTO batches(product_id,quantity,cost,expires_on,saleable) VALUES(?,'0','3.123456','2026-10-01',1)", (self.product['id'],))
        result = self.sync.sync_once()
        rows = self.rows(result['folder'], 'inventory')
        self.assertEqual(result['batch_count'], 2)
        self.assertEqual(rows[1][2:4], ('0001', '=商品'))
        self.assertEqual(rows[1][11], 0)
        self.assertEqual(rows[1][5], 3.123456)
        self.assertEqual(rows[1][9], -1)
        self.assertIn('已過期', rows[1][10])
        self.assertIsNone(rows[2][11])
        self.assertIn('未盤點', rows[2][10])
        workbook = load_workbook(Path(result['folder'])/FILENAMES['inventory'])
        self.assertEqual(workbook.active['D2'].data_type, 's')
        self.assertEqual(workbook.active.freeze_panes, 'A2')
        workbook.close()

    def test_price_changes_reach_all_three_without_replacing_external_code(self):
        first = self.sync.sync_once()
        self.conn.execute("UPDATE products SET price='0'")
        second = self.sync.sync_once()
        self.assertNotEqual(first['folder'], second['folder'])
        self.assertEqual(self.rows(second['folder'], 'inventory')[1][6], 0)
        invoice = self.rows(second['folder'], 'invoice')[1]
        self.assertEqual(invoice[1], '000099')
        self.assertEqual(invoice[3:], (0, '包', 1, 0))
        catalog = load_workbook(Path(second['folder'])/FILENAMES['catalog'])
        self.assertEqual(catalog.active['E11'].value, 0)
        self.assertEqual(catalog.active['K11'].value, 80)
        catalog.close()

    def test_catalog_keeps_all_other_zip_parts_and_blank_price(self):
        self.conn.execute('UPDATE products SET price=NULL')
        path = self.root/'catalog.xlsx'
        write_catalog(self.conn, self.template, path)
        with ZipFile(self.template) as source, ZipFile(path) as output:
            self.assertEqual(source.namelist(), output.namelist())
            for name in source.namelist():
                if name != 'xl/worksheets/sheet1.xml':
                    self.assertEqual(source.read(name), output.read(name))
        book = load_workbook(path)
        self.assertIsNone(book.active['E11'].value)
        self.assertEqual(book.active['K11'].value, 80)
        book.close()

    def test_failed_sync_retains_last_complete_generation_and_can_retry(self):
        first = self.sync.sync_once()
        self.conn.execute("UPDATE products SET price='46'")
        with patch('app.excel_sync.write_catalog', side_effect=OSError('disk busy')):
            with self.assertRaises(OSError):
                self.sync.sync_once()
        self.assertEqual(self.sync.status()['last'], first)
        self.assertEqual(json.loads(self.sync.manifest.read_text('utf-8'))['folder'], first['folder'])
        with self.assertRaises(Problem):
            self.sync.download()
        self.assertEqual(len(list(self.sync.root.glob('*/snapshot.json'))), 1)
        self.sync.sync_once()
        self.assertIsNone(self.sync.status()['error'])
        self.assertTrue(self.sync.download().is_file())

    def test_unchanged_snapshot_is_reused_and_only_two_generations_retained(self):
        first = self.sync.sync_once()
        self.assertEqual(self.sync.sync_once(), first)
        for price in ('46', '47', '48'):
            self.conn.execute('UPDATE products SET price=?', (price,))
            self.sync.sync_once()
        self.assertEqual(len(list(self.sync.root.glob('*/snapshot.json'))), 2)

    def test_changed_template_product_is_rejected_without_publishing(self):
        first = self.sync.sync_once()
        workbook = load_workbook(self.template)
        workbook.active['E9'] = '不同商品'
        workbook.save(self.template)
        workbook.close()
        with self.assertRaisesRegex(ValueError, '來源品項已變更'):
            self.sync.sync_once()
        self.assertEqual(self.sync.status()['last'], first)

    def test_edit_during_build_remains_pending_for_next_snapshot(self):
        original = write_catalog
        def edit_then_write(*args):
            self.conn.execute("UPDATE products SET price='90'")
            self.sync.request()
            return original(*args)
        with patch('app.excel_sync.write_catalog', side_effect=edit_then_write):
            first = self.sync.sync_once()
        self.assertTrue(self.sync.status()['pending'])
        self.assertEqual(self.rows(first['folder'], 'invoice')[1][3], 45)
        second = self.sync.sync_once()
        self.assertFalse(self.sync.status()['pending'])
        self.assertEqual(self.rows(second['folder'], 'invoice')[1][3], 90)

    def test_authenticated_routes_and_committed_writes_schedule_sync(self):
        client = self.app.test_client()
        self.assertEqual(client.get('/api/inventory-export/download').status_code, 401)
        with client.session_transaction() as session:
            session['user'] = 1
            session['csrf'] = 'test'
        self.sync.sync_once()
        response = client.patch('/api/products/'+str(self.product['id']), json={'price':'75','expected_version':1}, headers={'X-CSRF-Token':'test'})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(self.sync.status()['pending'])
        self.assertEqual(client.get('/api/inventory-export/download').status_code, 409)
        self.sync.sync_once()
        with client.get('/api/inventory-export/download') as download:
            self.assertEqual(download.status_code, 200)
        response = client.patch('/api/products/'+str(self.product['id']), json={'price':'88','expected_version':1}, headers={'X-CSRF-Token':'test'})
        self.assertEqual(response.status_code, 409)
        self.assertFalse(self.sync.status()['pending'])
