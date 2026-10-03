import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

import openpyxl

from app.common import Problem
from app.db import connect_db, init_db
from app.exports import build_invoice_xlsx, export_state
from app.products import create_product, update_product


class ExportTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.conn = connect_db(self.root/'db')
        self.addCleanup(self.conn.close)
        init_db(self.conn)
        self.p = create_product(self.conn, {'code':'0001','name':'測試','unit':'包','price':'45'},1)
        self.conn.execute('''INSERT INTO invoice_items(source,category,code,name,price,unit,tax,kind,product_id)
                             VALUES('one','1734','OLD','原名','20','包',1,0,?)''',(self.p['id'],))

    def test_seven_columns_values_and_identifier(self):
        result = build_invoice_xlsx(self.conn, self.root/'exports')
        wb = openpyxl.load_workbook(self.root/'exports'/result['filename'], read_only=True)
        rows = list(wb.active.values)
        wb.close()
        self.assertEqual(len(rows[0]), 7)
        self.assertEqual(rows[1], ('1734','OLD','測試',45,'包',1,0))

    def test_linked_products_keep_external_invoice_identifiers(self):
        self.conn.execute("UPDATE invoice_items SET code='0000123'")
        update_product(self.conn,self.p['id'],{'code':'NEW-INTERNAL','price':'55'},1,1)
        result=build_invoice_xlsx(self.conn,self.root/'exports')
        wb=openpyxl.load_workbook(self.root/'exports'/result['filename'],read_only=True)
        row=list(wb.active.values)[1]
        wb.close()
        self.assertEqual(row[1],'0000123')
        self.assertEqual(row[3],55)

    def test_missing_price_blocks_export(self):
        update_product(self.conn,self.p['id'],{'price':None},1,1)
        with self.assertRaises(Problem):
            build_invoice_xlsx(self.conn,self.root/'exports')

    def test_unmatched_invoice_blocks_export(self):
        self.conn.execute('UPDATE invoice_items SET product_id=NULL')
        with self.assertRaises(Problem):
            build_invoice_xlsx(self.conn,self.root/'exports')

    def test_fixed_vegetable_tax_preserved(self):
        self.conn.execute("INSERT INTO invoice_items(source,category,code,name,price,unit,tax,kind,fixed) VALUES('two','1733','A00010','蔬菜10元','10','個',0,0,1)")
        result=build_invoice_xlsx(self.conn,self.root/'exports')
        wb=openpyxl.load_workbook(self.root/'exports'/result['filename'],read_only=True)
        self.assertEqual(list(wb.active.values)[2][-2:],(0,0))
        wb.close()

    def test_price_change_invalidates_export_snapshot(self):
        result=build_invoice_xlsx(self.conn,self.root/'exports')
        self.assertFalse(export_state(self.conn)['pending'])
        update_product(self.conn,self.p['id'],{'price':'50'},1,1)
        self.assertTrue(export_state(self.conn)['pending'])
        wb=openpyxl.load_workbook(self.root/'exports'/result['filename'],read_only=True)
        self.assertEqual(list(wb.active.values)[1][3],45)
        wb.close()

    def test_latest_export_marker_tracks_last_successful_export(self):
        from app.exports import build_invoice_xlsx
        with patch('app.exports.uuid4') as new_id:
            new_id.return_value.hex = 'ffffffffffffffffffffffffffffffff'
            first = build_invoice_xlsx(self.conn, self.root/'exports')
            new_id.return_value.hex = '00000000000000000000000000000000'
            last = build_invoice_xlsx(self.conn, self.root/'exports')
        marker = self.conn.execute("SELECT value FROM metadata WHERE key='latest_invoice_export'").fetchone()
        self.assertIsNotNone(marker)
        self.assertEqual(marker[0], last['export_id'])
